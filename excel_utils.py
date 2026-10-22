"""Lectura y reserva segura de bloques en el libro Excel compartido."""
from pathlib import Path
import unicodedata

from filelock import FileLock, Timeout
from openpyxl import load_workbook

LOCK_TIMEOUT_SECONDS = 20


def _normalizar(valor):
    texto = "" if valor is None else str(valor).strip().casefold()
    texto = "".join(
        caracter for caracter in unicodedata.normalize("NFD", texto)
        if unicodedata.category(caracter) != "Mn"
    )
    return " ".join(texto.split())


def _normalizar_bloque(valor):
    return _normalizar(valor).replace("–", "-").replace("—", "-").replace(" - ", "-").replace(" ", "")



def _localizar_celda(planilla, dia, bloque):
    filas = planilla.iter_rows()
    encabezado = next(filas, None)

    if not encabezado:
        raise ValueError("La hoja Excel no contiene encabezados.")

    columna_dia = next(
        (
            i for i, celda in enumerate(encabezado)
            if _normalizar(celda.value) == _normalizar(dia)
        ),
        None,
    )

    if columna_dia is None:
        raise ValueError(
            f"No se encontró el día '{dia}' en los encabezados de Excel."
        )

    for fila in filas:
        if (
            fila
            and _normalizar_bloque(fila[0].value)
            == _normalizar_bloque(bloque)
        ):
            if columna_dia >= len(fila):
                raise ValueError(
                    "La estructura de la hoja Excel no coincide "
                    "con el horario esperado."
                )
            return fila[columna_dia]

    raise ValueError(f"No se encontró el bloque horario '{bloque}' en Excel.")



def _bloqueo(archivo):
    # El bloqueo entre procesos evita que dos workers de Gunicorn reserven
    # simultáneamente la misma celda. Todos los accesos de esta app usan el mismo lock.
    ruta_lock = f"{Path(archivo).resolve()}.lock"
    return FileLock(ruta_lock, timeout=LOCK_TIMEOUT_SECONDS)


def horario_disponible(archivo, hoja, dia, bloque):
    """Devuelve True si la celda del día/bloque está vacía."""
    with _bloqueo(archivo):
        libro = load_workbook(archivo, read_only=True, data_only=True)
        try:
            if hoja not in libro.sheetnames:
                raise ValueError(f"No existe la hoja '{hoja}' en el libro Excel.")
            celda = _localizar_celda(libro[hoja], dia, bloque)
            return celda.value is None or not str(celda.value).strip()
        finally:
            libro.close()


def reservar_horario(archivo, hoja, dia, bloque, nombre_docente):
    """Comprueba y escribe la reserva bajo un único bloqueo; False significa ocupado."""
    with _bloqueo(archivo):
        libro = load_workbook(archivo)
        try:
            if hoja not in libro.sheetnames:
                raise ValueError(f"No existe la hoja '{hoja}' en el libro Excel.")
            celda = _localizar_celda(libro[hoja], dia, bloque)
            if celda.value is not None and str(celda.value).strip():
                return False
            celda.value = nombre_docente.strip()
            libro.save(archivo)
            return True
        finally:
            libro.close()



def obtener_ocupacion(archivo, hoja):
    """Entrega la matriz para el panel respetando las columnas reales de Excel."""
    with _bloqueo(archivo):
        libro = load_workbook(archivo, read_only=True, data_only=True)
        try:
            if hoja not in libro.sheetnames:
                raise ValueError(f"No existe la hoja '{hoja}' en el libro Excel.")

            filas = list(libro[hoja].iter_rows(values_only=True))
            if not filas:
                raise ValueError("La hoja Excel está vacía.")

            encabezado = filas[0]

            # Guardar el índice real de cada día, incluso si hay columnas vacías.
            columnas_dias = [
                (indice, str(valor).strip())
                for indice, valor in enumerate(encabezado)
                if indice > 0 and valor is not None and str(valor).strip()
            ]
            dias = [dia for _, dia in columnas_dias]

            horarios = []
            for fila in filas[1:]:
                if not fila or fila[0] is None or not str(fila[0]).strip():
                    continue

                bloque = str(fila[0]).strip()
                estados = []

                for indice, dia in columnas_dias:
                    valor = fila[indice] if indice < len(fila) else None
                    docente = "" if valor is None else str(valor).strip()
                    estados.append({
                        "dia": dia,
                        "ocupado": bool(docente),
                        "docente": docente,
                    })

                horarios.append({
                    "bloque": bloque,
                    "estados": estados,
                })

            return {"dias": dias, "horarios": horarios}
        finally:
            libro.close()


            
def liberar_horario(archivo, hoja, dia, bloque):
    """Libera un bloque ocupado bajo el mismo bloqueo que las reservas."""
    with _bloqueo(archivo):
        libro = load_workbook(archivo)
        try:
            if hoja not in libro.sheetnames:
                raise ValueError(f"No existe la hoja '{hoja}' en el libro Excel.")

            celda = _localizar_celda(libro[hoja], dia, bloque)

            if celda.value is None or not str(celda.value).strip():
                return False

            celda.value = None
            libro.save(archivo)
            return True
        finally:
            libro.close()

