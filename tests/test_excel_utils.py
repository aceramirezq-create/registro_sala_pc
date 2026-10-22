import tempfile
import unittest
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

from openpyxl import Workbook, load_workbook

from excel_utils import horario_disponible, obtener_ocupacion, reservar_horario


class ExcelReservationTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.path = str(Path(self.temp_dir.name) / "calendario.xlsx")
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Hoja1"
        sheet.append(["Bloque", "Lunes", "Martes", "Miércoles", "Jueves", "Viernes"])
        sheet.append(["08:30 - 10:00", None, "Docente existente", None, None, None])
        sheet.append(["10:15 - 11:45", None, None, None, None, None])
        workbook.save(self.path)
        workbook.close()

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_dia_acepta_sin_acento(self):
        self.assertTrue(horario_disponible(self.path, "Hoja1", "miercoles", "08:30 - 10:00"))

    def test_no_reserva_sobre_bloque_ocupado(self):
        self.assertFalse(reservar_horario(self.path, "Hoja1", "Martes", "08:30 - 10:00", "Nuevo docente"))

    def test_reserva_actualiza_panel(self):
        self.assertTrue(reservar_horario(self.path, "Hoja1", "Miércoles", "10:15 - 11:45", "Ana Docente"))
        estado = obtener_ocupacion(self.path, "Hoja1")
        bloque = next(x for x in estado["horarios"] if x["bloque"] == "10:15 - 11:45")
        miercoles = next(x for x in bloque["estados"] if x["dia"] == "Miércoles")
        self.assertTrue(miercoles["ocupado"])
        self.assertEqual(miercoles["docente"], "Ana Docente")

    def test_solo_una_reserva_concurrente_gana(self):
        def reservar(nombre):
            return reservar_horario(self.path, "Hoja1", "Lunes", "10:15 - 11:45", nombre)

        with ThreadPoolExecutor(max_workers=2) as executor:
            resultados = list(executor.map(reservar, ["Docente A", "Docente B"]))
        self.assertEqual(resultados.count(True), 1)
        self.assertEqual(resultados.count(False), 1)
        workbook = load_workbook(self.path, data_only=True)
        try:
            self.assertIn(workbook["Hoja1"]["B3"].value, {"Docente A", "Docente B"})
        finally:
            workbook.close()


if __name__ == "__main__":
    unittest.main()
